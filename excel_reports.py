"""
excel_reports.py — Génération des rapports Excel (.xlsx) du module caille,
via openpyxl (bibliothèque 100% Python, sans dépendance système — compatible
avec un environnement serverless comme Vercel).
"""
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

HEADER_FILL = PatternFill(start_color="6E4622", end_color="6E4622", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill(start_color="F0E8DB", end_color="F0E8DB", fill_type="solid")
TOTAL_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center")


def _nouveau_classeur(titre_feuille):
    wb = Workbook()
    ws = wb.active
    ws.title = titre_feuille[:31]
    return wb, ws


def _remplir_feuille(ws, colonnes, lignes, ligne_totaux=None):
    ws.append(colonnes)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER

    for ligne in lignes:
        ws.append(ligne)

    if ligne_totaux is not None:
        ws.append(ligne_totaux)
        for cell in ws[ws.max_row]:
            cell.fill = TOTAL_FILL
            cell.font = TOTAL_FONT

    for col in ws.columns:
        largeur = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(largeur, 42)


def _vers_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def excel_ponte(entrees, totaux):
    wb, ws = _nouveau_classeur("Ponte")
    lignes = [[e.date_jour, e.lot.nom, e.oeufs_pondus, e.oeufs_vendus, e.prix_unitaire_vente,
               e.montant_vente, e.oeufs_casses, e.oeufs_autoconsommes] for e in entrees]
    total = ["", "TOTAL", totaux["pondus"], totaux["vendus"], "", totaux["revenu"], totaux["casses"], ""]
    _remplir_feuille(ws, ["Date", "Lot", "Pondus", "Vendus", "Prix/u (F)", "Montant (F)", "Cassés", "Autoconso."], lignes, total)
    return _vers_bytes(wb)


def excel_mortalite(entrees, totaux):
    wb, ws = _nouveau_classeur("Mortalité")
    lignes = [[e.date_jour, e.lot.nom, e.cailles_mortes, e.cailletons_morts, e.cause or "", e.notes or ""] for e in entrees]
    total = ["", "TOTAL", totaux["cailles"], totaux["cailletons"], "", ""]
    _remplir_feuille(ws, ["Date", "Lot", "Cailles mortes", "Cailletons morts", "Cause", "Notes"], lignes, total)
    return _vers_bytes(wb)


def excel_naissances(entrees, total):
    wb, ws = _nouveau_classeur("Naissances")
    lignes = [[e.date_jour, e.lot.nom, e.nombre_cailletons,
               "Éclosion" if e.origine == "eclosion" else "Achat", e.notes or ""] for e in entrees]
    total_row = ["", "TOTAL", total, "", ""]
    _remplir_feuille(ws, ["Date", "Lot", "Nombre", "Origine", "Notes"], lignes, total_row)
    return _vers_bytes(wb)


def excel_achats(achats, totaux):
    wb, ws = _nouveau_classeur("Achats provende")
    lignes = [[a.date_achat, a.type_provende.nom, a.fournisseur.nom if a.fournisseur else "—",
               a.quantite_kg, a.prix_total, a.prix_unitaire_kg] for a in achats]
    total = ["", "TOTAL", "", totaux["kg"], totaux["cout"], ""]
    _remplir_feuille(ws, ["Date", "Type", "Fournisseur", "Quantité (kg)", "Prix total (F)", "Prix/kg (F)"], lignes, total)
    return _vers_bytes(wb)


def excel_consommation(entrees, total_kg):
    wb, ws = _nouveau_classeur("Consommation")
    lignes = [[c.date_jour, c.lot.nom if c.lot else "Tous lots", c.type_provende.nom, c.quantite_kg, c.notes or ""] for c in entrees]
    total = ["", "", "TOTAL", total_kg, ""]
    _remplir_feuille(ws, ["Date", "Lot", "Type", "Quantité (kg)", "Notes"], lignes, total)
    return _vers_bytes(wb)


def excel_taches(taches):
    wb, ws = _nouveau_classeur("Cahier de charges")
    libelle = {"a_faire": "À faire", "en_cours": "En cours", "fait": "Fait"}
    lignes = [[t.titre, t.categorie, t.lot.nom if t.lot else "—",
               t.date_prevue.strftime("%d/%m/%Y") if t.date_prevue else "—",
               t.recurrence or "aucune", libelle.get(t.statut, t.statut)] for t in taches]
    _remplir_feuille(ws, ["Titre", "Catégorie", "Lot", "Date prévue", "Récurrence", "Statut"], lignes)
    return _vers_bytes(wb)


def excel_fournisseurs(fournisseurs):
    wb, ws = _nouveau_classeur("Fournisseurs")
    lignes = [[f.nom, f.contact or "—", f.adresse or "—", f.notes or ""] for f in fournisseurs]
    _remplir_feuille(ws, ["Nom", "Contact", "Adresse", "Notes"], lignes)
    return _vers_bytes(wb)


def excel_lot(lot, pontes, mortalites, naissances):
    wb = Workbook()
    ws_info = wb.active
    ws_info.title = "Résumé"
    infos = [
        ["Nom", lot.nom], ["Type", lot.type_lot],
        ["Mise en place", lot.date_mise_en_place], ["Effectif initial", lot.effectif_initial],
        ["Effectif actuel", lot.effectif_actuel], ["Statut", lot.statut],
        ["Œufs pondus (total)", lot.total_oeufs_pondus], ["Revenu œufs (total, F)", lot.total_revenu_oeufs],
        ["Mortalité cailles (total)", lot.total_mortalite_cailles],
        ["Mortalité cailletons (total)", lot.total_mortalite_cailletons],
        ["Provende consommée (kg)", lot.total_consommation_kg],
        ["Indice de consommation (kg/œuf)", lot.indice_consommation if lot.indice_consommation is not None else "—"],
        ["Coût provende estimé (F)", lot.cout_provende_estime],
        ["Marge estimée (F)", lot.marge_estimee],
    ]
    for ligne in infos:
        ws_info.append(ligne)
    for row in ws_info.iter_rows(min_col=1, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)
    for col in ws_info.columns:
        largeur = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws_info.column_dimensions[col[0].column_letter].width = min(largeur, 42)

    ws_p = wb.create_sheet("Ponte")
    _remplir_feuille(ws_p, ["Date", "Pondus", "Vendus", "Montant (F)"],
                      [[p.date_jour, p.oeufs_pondus, p.oeufs_vendus, p.montant_vente] for p in pontes])

    ws_m = wb.create_sheet("Mortalité")
    _remplir_feuille(ws_m, ["Date", "Cailles", "Cailletons", "Cause"],
                      [[m.date_jour, m.cailles_mortes, m.cailletons_morts, m.cause or ""] for m in mortalites])

    ws_n = wb.create_sheet("Naissances")
    _remplir_feuille(ws_n, ["Date", "Nombre", "Origine"],
                      [[n.date_jour, n.nombre_cailletons, n.origine] for n in naissances])

    return _vers_bytes(wb)


def excel_rentabilite(lots, stats_mensuelles):
    wb = Workbook()
    ws_lots = wb.active
    ws_lots.title = "Par lot"
    lignes = [[l.nom, l.total_oeufs_pondus, l.total_revenu_oeufs, l.total_consommation_kg,
               l.indice_consommation if l.indice_consommation is not None else "—",
               l.cout_provende_estime, l.total_depenses, l.marge_estimee] for l in lots]
    _remplir_feuille(ws_lots, ["Lot", "Œufs pondus", "Revenu (F)", "Provende consommée (kg)",
                                "Indice conso. (kg/œuf)", "Coût provende estimé (F)",
                                "Autres dépenses (F)", "Marge estimée (F)"], lignes)

    ws_mois = wb.create_sheet("Par mois")
    lignes_mois = [[s["mois"], s["revenu"], s["cout_provende"], s["depenses"], s["marge"]] for s in stats_mensuelles]
    _remplir_feuille(ws_mois, ["Mois", "Revenu œufs (F)", "Coût provende achetée (F)",
                                "Autres dépenses (F)", "Marge (F)"], lignes_mois)

    return _vers_bytes(wb)


def excel_depenses(entrees, total):
    wb, ws = _nouveau_classeur("Dépenses")
    lignes = [[d.date_depense, d.categorie, d.lot.nom if d.lot else "—", d.montant, d.notes or ""] for d in entrees]
    total_row = ["", "", "TOTAL", total, ""]
    _remplir_feuille(ws, ["Date", "Catégorie", "Lot", "Montant (F)", "Notes"], lignes, total_row)
    return _vers_bytes(wb)